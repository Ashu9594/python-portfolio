#!/usr/bin/env python3
"""
Build styled DOCX + PDF from resume/Ashutosh_Aagawane_Senior_Backend_Lead_ATS.txt

Outputs (same folder):
  - Ashutosh_Aagawane_Resume_Formatted.docx
  - Ashutosh_Aagawane_Resume_Formatted.pdf

Regenerate from repo root:
  python3 -m venv .venv-resume
  .venv-resume/bin/pip install -r resume/requirements.txt
  .venv-resume/bin/python resume/build_formatted_resume.py

To attach the PDF to the portfolio site, copy it to ../resume.pdf (next to index.html).
"""

from __future__ import annotations

import html
import sys
from pathlib import Path

# --- Paths ---
ROOT = Path(__file__).resolve().parents[1]
RESUME_DIR = Path(__file__).resolve().parent
SOURCE_TXT = RESUME_DIR / "Ashutosh_Aagawane_Senior_Backend_Lead_ATS.txt"
OUT_DOCX = RESUME_DIR / "Ashutosh_Aagawane_Resume_Formatted.docx"
OUT_PDF = RESUME_DIR / "Ashutosh_Aagawane_Resume_Formatted.pdf"

# --- Brand colors (teal / slate) ---
RGB_TEAL = (14, 116, 110)
RGB_SLATE = (30, 41, 59)
RGB_MUTED = (71, 85, 105)


def parse_sections(text: str) -> tuple[list[str], list[tuple[str, str]]]:
    lines = text.splitlines()
    preamble: list[str] = []
    i = 0
    while i < len(lines):
        s = lines[i].strip()
        if s.startswith("=") and len(s) > 40:
            break
        preamble.append(lines[i])
        i += 1

    sections: list[tuple[str, str]] = []
    while i < len(lines):
        s = lines[i].strip()
        if s.startswith("=") and len(s) > 40:
            i += 1
            if i >= len(lines):
                break
            title = lines[i].strip()
            i += 1
            if i < len(lines) and lines[i].strip().startswith("="):
                i += 1
            body_lines: list[str] = []
            while i < len(lines):
                s2 = lines[i].strip()
                if s2.startswith("=") and len(s2) > 40:
                    break
                body_lines.append(lines[i])
                i += 1
            body = "\n".join(body_lines).strip()
            sections.append((title, body))
        else:
            i += 1

    return preamble, sections


def should_skip_section(title: str) -> bool:
    t = title.upper()
    return "HOW TO USE" in t


def clean_body(body: str) -> str:
    lines = []
    for line in body.splitlines():
        if "Estimated readiness:" in line:
            continue
        lines.append(line)
    return "\n".join(lines).strip()


def split_experience(body: str) -> list[dict]:
    jobs: list[dict] = []
    title: str | None = None
    dates: str | None = None
    bullets: list[str] = []

    def flush():
        nonlocal title, dates, bullets, jobs
        if title:
            jobs.append({"title": title, "dates": dates or "", "bullets": bullets[:]})
        title = None
        dates = None
        bullets = []

    for raw in body.splitlines():
        line = raw.strip()
        if not line:
            continue
        if line.startswith("•"):
            bullets.append(line[1:].strip())
            continue
        if "|" in line and ("—" in line or " – " in line or " - " in line):
            flush()
            title = line
            dates = None
            bullets = []
            continue
        if title and dates is None:
            dates = line
            continue
        if title:
            bullets.append(line)

    flush()
    return jobs


def escape_pdf(s: str) -> str:
    return html.escape(s, quote=False).replace("\n", "<br/>")


def build_docx(preamble: list[str], sections: list[tuple[str, str]]) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor, Inches

    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(0.65)
        s.bottom_margin = Inches(0.65)
        s.left_margin = Inches(0.75)
        s.right_margin = Inches(0.75)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(*RGB_SLATE)

    # --- Header ---
    name = preamble[0].strip() if preamble else "ASHUTOSH AAGAWANE"
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(name)
    r.bold = True
    r.font.size = Pt(22)
    r.font.color.rgb = RGBColor(*RGB_SLATE)
    r.font.name = "Calibri Light"

    target_lines = []
    contact_lines = []
    for section_title, body in sections:
        if section_title.upper().startswith("TARGET ROLES"):
            target_lines = [ln.strip() for ln in body.splitlines() if ln.strip()]
            break
    for line in preamble[1:]:
        if line.strip():
            contact_lines.append(line.strip())

    if target_lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(target_lines[0])
        r.italic = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(*RGB_TEAL)
        if len(target_lines) > 1:
            r = p.add_run("\n" + target_lines[1])
            r.italic = True
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(*RGB_MUTED)

    for cl in contact_lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(cl)
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(*RGB_MUTED)

    doc.add_paragraph()

    teal = RGBColor(*RGB_TEAL)
    slate = RGBColor(*RGB_SLATE)

    def add_heading(text: str) -> None:
        p = doc.add_paragraph()
        r = p.add_run(text.upper())
        r.bold = True
        r.font.size = Pt(11)
        r.font.color.rgb = teal
        r.font.name = "Calibri"
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True

    for title, body in sections:
        if should_skip_section(title):
            continue
        if title.upper().startswith("TARGET ROLES"):
            continue
        body = clean_body(body)
        if not body:
            continue

        t_up = title.upper()
        if t_up.startswith("PROFESSIONAL EXPERIENCE"):
            add_heading("Professional experience")
            for job in split_experience(body):
                p = doc.add_paragraph()
                r = p.add_run(job["title"])
                r.bold = True
                r.font.size = Pt(10.5)
                r.font.color.rgb = slate
                p.paragraph_format.space_before = Pt(8)
                p.paragraph_format.keep_with_next = True

                p = doc.add_paragraph()
                r = p.add_run(job["dates"])
                r.italic = True
                r.font.size = Pt(10)
                r.font.color.rgb = RGBColor(*RGB_MUTED)
                p.paragraph_format.space_after = Pt(3)

                for b in job["bullets"]:
                    bp = doc.add_paragraph(style="List Bullet")
                    bp.paragraph_format.left_indent = Inches(0.25)
                    bp.paragraph_format.space_after = Pt(2)
                    br = bp.add_run(b)
                    br.font.size = Pt(10.5)
            continue

        if t_up.startswith("SELECTED IMPACT"):
            add_heading("Selected impact & scale")
        elif t_up.startswith("CORE COMPETENCIES"):
            add_heading("Core competencies")
        elif t_up.startswith("PROFESSIONAL SUMMARY"):
            add_heading("Professional summary")
        elif t_up.startswith("EDUCATION"):
            add_heading("Education")
        else:
            add_heading(title.replace(" (ATS KEYWORDS)", "").replace(" (customize per application)", ""))

        if t_up.startswith("CORE COMPETENCIES"):
            for line in body.splitlines():
                line = line.strip()
                if not line:
                    continue
                bp = doc.add_paragraph(style="List Bullet")
                bp.paragraph_format.space_after = Pt(2)
                bp.add_run(line)
        elif t_up.startswith("SELECTED IMPACT"):
            for line in body.splitlines():
                line = line.strip()
                if not line:
                    continue
                if line.startswith("•"):
                    line = line[1:].strip()
                bp = doc.add_paragraph(style="List Bullet")
                bp.paragraph_format.space_after = Pt(2)
                bp.add_run(line)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            p.add_run(body)

    doc.save(OUT_DOCX)


def build_pdf(preamble: list[str], sections: list[tuple[str, str]]) -> None:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_JUSTIFY
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            name="ResumeName",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=20,
            textColor=colors.HexColor("#0f172a"),
            alignment=TA_CENTER,
            spaceAfter=4,
        )
    )
    styles.add(
        ParagraphStyle(
            name="ResumeTagline",
            parent=styles["Normal"],
            fontName="Helvetica-Oblique",
            fontSize=10,
            textColor=colors.HexColor("#0d9488"),
            alignment=TA_CENTER,
            spaceAfter=2,
        )
    )
    styles.add(
        ParagraphStyle(
            name="ResumeContact",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=8.5,
            textColor=colors.HexColor("#64748b"),
            alignment=TA_CENTER,
            spaceAfter=2,
        )
    )
    styles.add(
        ParagraphStyle(
            name="ResumeHeading",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=10.5,
            textColor=colors.HexColor("#0d9488"),
            alignment=TA_LEFT,
            spaceBefore=14,
            spaceAfter=6,
            borderPadding=0,
            borderWidth=0,
            borderColor=colors.HexColor("#0d9488"),
        )
    )
    styles.add(
        ParagraphStyle(
            name="ResumeBody",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=10,
            textColor=colors.HexColor("#1e293b"),
            alignment=TA_JUSTIFY,
            spaceAfter=6,
            leading=13,
        )
    )
    styles.add(
        ParagraphStyle(
            name="JobTitle",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=10.5,
            textColor=colors.HexColor("#0f172a"),
            alignment=TA_LEFT,
            spaceBefore=10,
            spaceAfter=2,
        )
    )
    styles.add(
        ParagraphStyle(
            name="JobDates",
            parent=styles["Normal"],
            fontName="Helvetica-Oblique",
            fontSize=9.5,
            textColor=colors.HexColor("#64748b"),
            alignment=TA_LEFT,
            spaceAfter=4,
        )
    )
    styles.add(
        ParagraphStyle(
            name="ResumeBullet",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=10,
            textColor=colors.HexColor("#334155"),
            alignment=TA_LEFT,
            leftIndent=16,
            firstLineIndent=-10,
            spaceAfter=4,
            leading=13,
        )
    )

    story: list = []
    doc = SimpleDocTemplate(
        str(OUT_PDF),
        pagesize=LETTER,
        leftMargin=0.72 * inch,
        rightMargin=0.72 * inch,
        topMargin=0.6 * inch,
        bottomMargin=0.6 * inch,
        title="Ashutosh Aagawane — Resume",
        author="Ashutosh Aagawane",
    )

    name = preamble[0].strip() if preamble else "ASHUTOSH AAGAWANE"
    story.append(Paragraph(escape_pdf(name), styles["ResumeName"]))

    target_lines = []
    for section_title, body in sections:
        if section_title.upper().startswith("TARGET ROLES"):
            target_lines = [ln.strip() for ln in body.splitlines() if ln.strip()]
            break
    if target_lines:
        story.append(Paragraph(escape_pdf(target_lines[0]), styles["ResumeTagline"]))
        if len(target_lines) > 1:
            story.append(Paragraph(escape_pdf(target_lines[1]), styles["ResumeContact"]))

    for line in preamble[1:]:
        if line.strip():
            story.append(Paragraph(escape_pdf(line.strip()), styles["ResumeContact"]))

    story.append(Spacer(1, 0.12 * inch))

    for title, body in sections:
        if should_skip_section(title):
            continue
        if title.upper().startswith("TARGET ROLES"):
            continue
        body = clean_body(body)
        if not body:
            continue

        t_up = title.upper()
        if t_up.startswith("PROFESSIONAL EXPERIENCE"):
            story.append(Paragraph("<b>PROFESSIONAL EXPERIENCE</b>", styles["ResumeHeading"]))
            for job in split_experience(body):
                story.append(Paragraph(escape_pdf(job["title"]), styles["JobTitle"]))
                story.append(Paragraph(escape_pdf(job["dates"]), styles["JobDates"]))
                for b in job["bullets"]:
                    story.append(Paragraph("\u2022 " + escape_pdf(b), styles["ResumeBullet"]))
            continue

        if t_up.startswith("SELECTED IMPACT"):
            story.append(Paragraph("<b>SELECTED IMPACT &amp; SCALE</b>", styles["ResumeHeading"]))
            bullets = []
            for line in body.splitlines():
                line = line.strip()
                if not line:
                    continue
                if line.startswith("•"):
                    line = line[1:].strip()
                bullets.append(line)
            for b in bullets:
                story.append(Paragraph("\u2022 " + escape_pdf(b), styles["ResumeBullet"]))
            continue

        if t_up.startswith("CORE COMPETENCIES"):
            story.append(Paragraph("<b>CORE COMPETENCIES</b>", styles["ResumeHeading"]))
            for line in body.splitlines():
                line = line.strip()
                if not line:
                    continue
                story.append(Paragraph("\u2022 " + escape_pdf(line), styles["ResumeBullet"]))
            continue

        if t_up.startswith("PROFESSIONAL SUMMARY"):
            story.append(Paragraph("<b>PROFESSIONAL SUMMARY</b>", styles["ResumeHeading"]))
            story.append(Paragraph(escape_pdf(body), styles["ResumeBody"]))
            continue

        if t_up.startswith("EDUCATION"):
            story.append(Paragraph("<b>EDUCATION</b>", styles["ResumeHeading"]))
            story.append(Paragraph(escape_pdf(body), styles["ResumeBody"]))
            continue

        story.append(Paragraph("<b>" + escape_pdf(title.upper()) + "</b>", styles["ResumeHeading"]))
        story.append(Paragraph(escape_pdf(body), styles["ResumeBody"]))

    doc.build(story)


def main() -> int:
    if not SOURCE_TXT.is_file():
        print(f"Missing source file: {SOURCE_TXT}", file=sys.stderr)
        return 1
    text = SOURCE_TXT.read_text(encoding="utf-8")
    preamble, sections = parse_sections(text)

    build_docx(preamble, sections)
    print(f"Wrote {OUT_DOCX}")

    build_pdf(preamble, sections)
    print(f"Wrote {OUT_PDF}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
